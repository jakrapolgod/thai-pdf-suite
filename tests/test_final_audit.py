# -*- coding: utf-8 -*-
"""
ตรวจสอบขั้นสุดท้าย - ห้ามเพี้ยนเด็ดขาด

ต่างจากชุดทดสอบอื่นตรงที่ **ไม่ใช้ PDF ที่เราสร้างเอง** แต่ให้ Word และ Excel ตัวจริง
เป็นคนสร้าง PDF ให้ เพราะนั่นคือที่มาจริงของเอกสารไทยที่ผู้ใช้จะเอามาแปลง

ครอบคลุม
  1. ข้อความไทยที่ถูกต้องอยู่แล้ว ต้องไม่ถูกแตะแม้แต่ตัวเดียว (กันตัวซ่อมทำพัง)
  2. จัดระเบียบซ้ำต้องได้ผลเดิม และห้ามทำพยัญชนะหาย
  3. Word -> PDF -> แอป  เทียบทีละ code point
  4. Excel -> PDF -> แอป  เทียบทีละ code point
  5. วนครบวง PDF -> แอป -> Word -> PDF  ข้อความต้องเหมือนเดิม

รันด้วย:  python tests/test_final_audit.py
"""
import os
import sys
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import fitz                                        # noqa: E402
from core import thaitext as T                     # noqa: E402
from core.build_docx import build_docx             # noqa: E402
from core.extract import extract_document          # noqa: E402

OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
PASS, FAIL = 0, 0
FAILURES = []


def cp(s: str, limit: int = 40) -> str:
    out = " ".join(f"{ord(c):04X}" for c in s[:limit])
    return out + (" …" if len(s) > limit else "")


def check(name, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1
    else:
        FAIL += 1
        FAILURES.append(f"{name}  {extra}")
        print(f"  [ตก ] {name}  {extra}")


# ==================================================================== คลังคำ
# ครอบคลุมทุกวรรณยุกต์ ทุกสระ ทัณฑฆาต ไม้ไต่คู้ ไม้ยมก เลขไทย และสัญลักษณ์
CORPUS = [
    # วรรณยุกต์เอก
    "ก่อน", "ต่อ", "ผ่าน", "ใหม่", "อยู่", "ที่", "เพื่อ", "ค่า", "น่า", "แต่",
    # วรรณยุกต์โท
    "น้ำ", "ให้", "ได้", "ไม้", "ขึ้น", "ครั้ง", "ทั้ง", "นี้", "ผู้", "เก้า",
    # วรรณยุกต์ตรี
    "โต๊ะ", "ก๊าซ", "ตุ๊กตา", "จ๊ะ", "เจี๊ยบ", "ป๊า",
    # วรรณยุกต์จัตวา
    "เดี๋ยว", "จ๋า", "ป๋า", "ตี๋", "เก๋",
    # ทัณฑฆาต
    "ปาล์ม", "สิงห์", "จันทร์", "ศาสตร์", "เล่ห์", "โทรศัพท์", "อินเทอร์เน็ต",
    "สัมพันธ์", "บริสุทธิ์", "ทรัพย์สิน", "ลักษณ์",
    # ไม้ไต่คู้
    "เก็บ", "ก็", "เล็ก", "เห็น", "เป็น", "เด็ก", "เผ็ด",
    # สระอำ ทั้งมีและไม่มีวรรณยุกต์
    "ทำ", "จำ", "ดำ", "คำ", "น้ำ", "ค่ำ", "ซ้ำ", "ย้ำ", "ส้ม", "ทำงาน",
    # ฤ ฦ
    "ฤดู", "ฤทธิ์", "พฤษภาคม", "อังกฤษ", "พฤหัสบดี",
    # สระประสมและตัวสะกดยาก
    "เกี๊ยว", "เรียบ", "เหมือน", "เกือบ", "ด้วย", "เดี่ยว", "เปลี่ยน",
    "ปรารถนา", "ขอบคุณ", "สรรพสิ่ง", "พุทธศักราช", "กิจกรรม", "อุตสาหกรรม",
    "ประสิทธิภาพ", "วิเคราะห์", "สถาปัตยกรรม", "โครงสร้าง",
    # ไม้ยมก สัญลักษณ์ เลขไทย
    "ต่างๆ", "ฯลฯ", "กรุงเทพฯ", "๑๒๓๔๕๖๗๘๙๐", "๒๕๖๙", "฿",
    # ประโยคเต็ม
    "บริษัท Thai PDF Suite จำกัด ดำเนินกิจการมากว่าสามสิบปี",
    "ยอดขายน้ำมันปาล์มดิบรวมทั้งสิ้น ๑,๒๓๔ ตัน เพิ่มขึ้นจากปีก่อนร้อยละ ๑๒",
    "ผู้บริหารได้ตั้งเป้าหมายไว้ว่าจะรักษาอัตราการเติบโตนี้ให้ต่อเนื่อง",
    "ใบกำกับภาษี เลขที่ INV-2569/0012 วันที่ ๓๐ กันยายน พ.ศ. ๒๕๖๙",
    "สัญญาฉบับนี้ทำขึ้นระหว่างผู้ว่าจ้างและผู้รับจ้าง ณ วันที่ระบุข้างต้น",
    "เงื่อนไขการชำระเงิน ๓๐ วันนับจากวันที่ได้รับใบแจ้งหนี้ฉบับสมบูรณ์",
    "คุณภาพน้ำมันต้องผ่านเกณฑ์ที่กำหนด มิฉะนั้นถือว่าผิดสัญญา",
    "โปรดตรวจสอบความถูกต้องก่อนลงลายมือชื่อรับรองเอกสารทุกครั้ง",
]


print("\n=== 1. ข้อความไทยที่ถูกอยู่แล้ว ต้องไม่ถูกแตะ ===")
for item in CORPUS:
    got, issues = T.normalize(item)
    check(f"คงเดิม: {item[:30]}", got == item,
          f"\n         เดิม {cp(item)}\n         ได้  {cp(got)}")
    check(f"ไม่มีการแก้: {item[:24]}", not issues,
          f"แก้ไป {[i.kind for i in issues]}")
    check(f"audit สะอาด: {item[:24]}", not T.audit(item))
print(f"  ตรวจ {len(CORPUS)} รายการ")

print("\n=== 2. จัดระเบียบซ้ำต้องได้ผลเดิม และห้ามทำอักขระหาย ===")
SICK = [
    "นํ้า",              # นํ้า
    "เเละ",              # เเละ
    "ก้ับ",              # ก+โท+หันอากาศ+บ
    "ท้ิ้ง",        # ท+โท+อิ+โท+ง
    "นำ้",                    # นำ+โท
]
for item in CORPUS + SICK:
    once = T.normalize(item)[0]
    twice = T.normalize(once)[0]
    check(f"ทำซ้ำได้ผลเดิม: {item[:20]}", once == twice,
          f"\n         ครั้งแรก {cp(once)}\n         ครั้งสอง {cp(twice)}")
    # พยัญชนะและสระที่กินที่ ห้ามหายแม้แต่ตัวเดียว
    def bases(s):
        return [c for c in s if c not in T.COMBINING]
    check(f"อักขระฐานไม่หาย: {item[:20]}", bases(once) == bases(item) or item in SICK,
          f"\n         เดิม {bases(item)}\n         ได้  {bases(once)}")

print("\n=== 3. ข้อความทุกตัวต้องอยู่ในรูป NFC ===")
for item in CORPUS:
    check(f"NFC: {item[:24]}", unicodedata.normalize("NFC", item) == item)


# ==================================================================== Word จริง
def word_to_pdf(docx_path: str, pdf_path: str):
    import win32com.client as w
    if os.path.exists(pdf_path):
        os.remove(pdf_path)
    app = w.DispatchEx("Word.Application")
    app.Visible = False
    try:
        doc = app.Documents.Open(os.path.abspath(docx_path), ReadOnly=True)
        doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17)
        doc.Close(False)
    finally:
        app.Quit()


def excel_to_pdf(xlsx_path: str, pdf_path: str):
    import win32com.client as w
    if os.path.exists(pdf_path):
        os.remove(pdf_path)
    app = w.DispatchEx("Excel.Application")
    app.Visible = False
    app.DisplayAlerts = False
    try:
        wb = app.Workbooks.Open(os.path.abspath(xlsx_path), ReadOnly=True)
        wb.ExportAsFixedFormat(0, os.path.abspath(pdf_path))
        wb.Close(False)
    finally:
        app.Quit()


HAVE_OFFICE = True
try:
    import win32com.client  # noqa: F401
except ImportError:
    HAVE_OFFICE = False

if HAVE_OFFICE:
    print("\n=== 4. Word ตัวจริงสร้าง PDF แล้วให้แอปอ่านกลับ ===")
    from docx import Document as Docx                    # noqa: E402
    from docx.shared import Pt                           # noqa: E402
    from docx.oxml.ns import qn                          # noqa: E402

    PARAS = CORPUS[-8:] + [
        "รายการสินค้า น้ำมันปาล์มดิบ จำนวน ๑,๒๓๔ ตัน มูลค่า ๔๕,๖๗๘,๙๐๐ บาท",
        "ข้อความยาวเพื่อทดสอบการตัดบรรทัดกลางคำ " + "ประสิทธิภาพการบริหารจัดการต้นทุน" * 4,
    ]

    for font in ["TH Sarabun New", "Angsana New", "Tahoma"]:
        d = Docx()
        style = d.styles["Normal"]
        style.font.size = Pt(14)
        rpr = style.element.get_or_add_rPr()
        rf = rpr.makeelement(qn("w:rFonts"), {})
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rf.set(qn(a), font)
        rpr.insert(0, rf)
        for p in PARAS:
            d.add_paragraph(p)
        src = OUT / f"_wt_{abs(hash(font))}.docx"
        pdf = OUT / f"_wt_{abs(hash(font))}.pdf"
        d.save(str(src))
        word_to_pdf(str(src), str(pdf))

        doc = extract_document(str(pdf))
        got_paras = [p.text for page in doc.pages for p in page.paragraphs]

        for i, want in enumerate(PARAS):
            got = got_paras[i] if i < len(got_paras) else ""
            ok = got == want
            check(f"[{font}] ย่อหน้า {i + 1}", ok,
                  "" if ok else
                  f"\n         ต้องการ {want[:60]!r}\n         ได้     {got[:60]!r}"
                  f"\n         cp ต้องการ {cp(want, 25)}\n         cp ได้     {cp(got, 25)}")
        check(f"[{font}] จำนวนย่อหน้าครบ", len(got_paras) == len(PARAS),
              f"ได้ {len(got_paras)} จาก {len(PARAS)}")
        check(f"[{font}] โครงสร้างอักขระไทยสะอาด", not T.audit(doc.text))
        os.remove(src)
        os.remove(pdf)

    print("\n=== 5. Excel ตัวจริงสร้าง PDF แล้วให้แอปอ่านกลับ ===")
    from openpyxl import Workbook                        # noqa: E402
    from openpyxl.styles import Font as XFont            # noqa: E402

    ROWS = [
        ["รายการ", "จำนวน", "หน่วย", "มูลค่า"],
        ["น้ำมันปาล์มดิบ", "1,234", "ตัน", "45,678,900.00"],
        ["เมล็ดในปาล์ม", "567", "ตัน", "12,345,600.00"],
        ["กะลาปาล์ม ชั้นหนึ่ง", "89", "ตัน", "1,234,500.00"],
        ["ทะลายเปล่าอัดก้อน", "456", "ตัน", "567,800.00"],
        ["รวมทั้งสิ้น", "2,346", "ตัน", "59,826,800.00"],
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "รายงาน"
    for r in ROWS:
        ws.append(r)
    for row in ws.iter_rows():
        for c in row:
            c.font = XFont(name="TH Sarabun New", size=16)
    for col, width in zip("ABCD", (26, 12, 10, 20)):
        ws.column_dimensions[col].width = width
    xsrc = OUT / "_xt.xlsx"
    xpdf = OUT / "_xt.pdf"
    wb.save(str(xsrc))
    wb.close()
    excel_to_pdf(str(xsrc), str(xpdf))

    xdoc = extract_document(str(xpdf))
    flat = " ".join(xdoc.text.split())
    for row in ROWS:
        for cell in row:
            check(f"[Excel] พบค่า {cell[:24]}", cell in flat,
                  f"\n         cp ต้องการ {cp(cell)}")
    check("[Excel] โครงสร้างอักขระไทยสะอาด", not T.audit(xdoc.text))

    print("\n=== 6. วนครบวง: PDF -> แอป -> Word -> PDF ===")
    # ถ้าข้อความรอดครบหลังผ่าน Word จริงไปกลับ แปลว่าไม่มีจุดไหนทำให้เพี้ยน
    src_doc = extract_document(str(ROOT / "tests" / "sample_thai.pdf"))
    build_docx(src_doc, str(OUT / "_rt.docx"), font_name="TH Sarabun New")
    word_to_pdf(str(OUT / "_rt.docx"), str(OUT / "_rt.pdf"))
    back = extract_document(str(OUT / "_rt.pdf"))

    a = " ".join(src_doc.text.split())
    b = " ".join(back.text.split())
    check("ข้อความรอดครบหลังวนผ่าน Word", a == b,
          f"\n         ก่อน {len(a)} ตัว  หลัง {len(b)} ตัว"
          f"\n         ก่อน {a[:70]!r}\n         หลัง {b[:70]!r}")
    check("โครงสร้างอักขระไทยสะอาดหลังวนกลับ", not T.audit(back.text))

    # เก็บภาพไว้ดูด้วยตา
    fitz.open(str(OUT / "_rt.pdf"))[0].get_pixmap(dpi=140).save(
        str(OUT / "ตรวจสอบครั้งสุดท้าย.png"))
    os.remove(xsrc)
    os.remove(xpdf)
else:
    print("\n(ข้ามข้อ 4-6 เพราะไม่มี pywin32 สำหรับเรียก Word/Excel)")

print("\n" + "=" * 62)
if FAIL:
    print(f"ผ่าน {PASS}   ตก {FAIL}")
    print("\nรายการที่ตก:")
    for f in FAILURES[:20]:
        print("  -", f.split("\n")[0])
else:
    print(f"ผ่านทั้งหมด {PASS} รายการ ไม่พบการเพี้ยนแม้แต่จุดเดียว")
print("=" * 62)
sys.exit(1 if FAIL else 0)
