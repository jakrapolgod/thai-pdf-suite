# -*- coding: utf-8 -*-
"""
ทดสอบเส้นทางทำงานจริงตั้งแต่ PDF จนถึงไฟล์ Word/Excel
รวมถึงกรณี PDF ที่ฝังอักขระไทยแบบเสีย (มรดกฟอนต์เก่า) เพื่อพิสูจน์ว่าแอปซ่อมให้จริง

รันด้วย:  python tests/test_pipeline.py
"""
import io
import sys
from pathlib import Path

import fitz
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from core import thaitext as T                     # noqa: E402
from core import verify as V                       # noqa: E402
from core.build_docx import build_docx             # noqa: E402
from core.build_xlsx import build_xlsx             # noqa: E402
from core.extract import extract_document          # noqa: E402
from core.images_to_pdf import images_to_pdf       # noqa: E402

OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
FONT = "C:/Windows/Fonts/tahoma.ttf"

PASS, FAIL = 0, 0


def check(name, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  [ผ่าน] {name}")
    else:
        FAIL += 1
        print(f"  [ตก ] {name}  {extra}")


# ข้อความเสียที่จะฝังลงไฟล์ คู่กับรูปที่ถูกต้องซึ่งคาดหวังหลังแปลง
SICK_PAIRS = [
    ("\u0E19\u0E4D\u0E49\u0E32\u0E43\u0E08",              # นํ้าใจ แบบนิคหิต
     "\u0E19\u0E49\u0E33\u0E43\u0E08"),
    ("\u0E40\u0E40\u0E25\u0E30\u0E17\u0E35\u0E48",        # เเละที่
     "\u0E41\u0E25\u0E30\u0E17\u0E35\u0E48"),
    ("\u0E17\u0E49\u0E34\u0E49\u0E07\u0E02\u0E22\u0E30",  # ทิ้งขยะ ลำดับสลับ+ซ้ำ
     "\u0E17\u0E34\u0E49\u0E07\u0E02\u0E22\u0E30"),
    ("\u0E01\u0E49\u0E31\u0E1A\u0E02\u0E49\u0E32\u0E27",  # กั้บข้าว ลำดับสลับ
     "\u0E01\u0E31\u0E49\u0E1A\u0E02\u0E49\u0E32\u0E27"),
]


def make_sick_pdf(path: Path):
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_font(fontname="TH", fontfile=FONT)
    y = 100
    for bad, _ in SICK_PAIRS:
        page.insert_text((72, y), bad, fontname="TH", fontsize=16)
        y += 34
    doc.save(str(path))
    doc.close()


print("\n=== 1. PDF ที่มีอักขระไทยเสีย ต้องถูกซ่อมระหว่างอ่าน ===")
sick_pdf = OUT / "_sick.pdf"
make_sick_pdf(sick_pdf)
doc = extract_document(str(sick_pdf))
got = doc.text
for bad, good in SICK_PAIRS:
    check(f"ซ่อม {bad!r} เป็น {good!r}", good in got,
          f"ได้: {got!r}")
check("ไม่มีอักขระเสียหลงเหลือ", all(bad not in got for bad, _ in SICK_PAIRS))
check("ผ่านการตรวจโครงสร้างอักขระไทย", len(T.audit(got)) == 0)
check("มีบันทึกว่าแก้อะไรไปบ้าง", len(doc.issues) > 0,
      f"issues={len(doc.issues)}")

print("\n=== 2. เอกสารปกติ แปลงเป็น Word แล้วต้องตรงทุกตัวอักษร ===")
sample = ROOT / "tests" / "sample_thai.pdf"
d = extract_document(str(sample))
build_docx(d, str(OUT / "_p.docx"))
rw = V.verify_docx(d, str(OUT / "_p.docx"))
check("ตรงทุกชิ้น", rw.exact, f"{rw.items_matched}/{rw.items_total}")
check("ไม่มีปัญหาอักขระไทยในไฟล์ผลลัพธ์", not rw.thai_problems)
check("จำนวนตัวอักษรไม่หาย", rw.chars_output >= rw.chars_source,
      f"{rw.chars_source} -> {rw.chars_output}")

print("\n=== 3. เอกสารปกติ แปลงเป็น Excel แล้วต้องตรงทุกค่า ===")
build_xlsx(d, str(OUT / "_p.xlsx"))
rx = V.verify_xlsx(d, str(OUT / "_p.xlsx"))
check("ตรงทุกค่า", rx.exact, f"{rx.items_matched}/{rx.items_total}")
check("ไม่มีปัญหาอักขระไทยในไฟล์ผลลัพธ์", not rx.thai_problems)

print("\n=== 4. ไฟล์ Word ต้องตั้งฟอนต์ฝั่ง complex script ให้ครบ ===")
import zipfile                                          # noqa: E402
with zipfile.ZipFile(OUT / "_p.docx") as z:
    xml = z.read("word/document.xml").decode("utf-8")
check("มี w:cs (ฟอนต์ complex script)", 'w:cs="' in xml)
check("มี w:szCs (ขนาด complex script)", "w:szCs" in xml)
check("ประกาศภาษาไทย", 'w:bidi="th-TH"' in xml)

print("\n=== 5. ข้อความในไฟล์ Word ต้องตรงระดับ code point ===")
from core.build_docx import read_docx_text               # noqa: E402
docx_text = read_docx_text(str(OUT / "_p.docx"))
for phrase in ["Thai PDF Suite", "รายงานผลการดำเนินงานประจำไตรมาส",
               "น้ำมันปาล์มดิบ", "๑,๒๓๔"]:
    check(f"พบวลี {phrase}", phrase in docx_text)

print("\n=== 6. ตารางต้องถูกดึงครบทุกช่อง ===")
tables = [t for p in d.pages for t in p.tables]
check("พบตาราง 1 ตาราง", len(tables) == 1, f"พบ {len(tables)}")
if tables:
    rows = tables[0].rows
    check("จำนวนแถวครบ 6", len(rows) == 6, f"พบ {len(rows)}")
    check("หัวตารางถูกต้อง", rows[0][:2] == ["รายการ", "จำนวน"], f"{rows[0]}")
    check("แถวสุดท้ายถูกต้อง", rows[-1][-1] == "59,826,800.00", f"{rows[-1]}")

print("\n=== 7. Excel ต้องแปลงตัวเลขเป็นค่าตัวเลขจริง ===")
from openpyxl import load_workbook                       # noqa: E402
wb = load_workbook(OUT / "_p.xlsx")
ws = wb[[s for s in wb.sheetnames if "ตาราง" in s][0]]
vals = {ws.cell(r, 4).value for r in range(2, 7)}
check("มูลค่าเป็นตัวเลขจริง", all(isinstance(v, (int, float)) for v in vals), f"{vals}")
check("ข้อความไทยยังเป็นข้อความ", isinstance(ws.cell(2, 1).value, str))
check("ค่าตรงกับต้นฉบับ", 45678900.0 in vals, f"{vals}")
wb.close()

print("\n=== 8. รวมรูปเป็น PDF ===")
imgs = []
for i, size in enumerate([(1200, 900), (900, 1200), (600, 600)]):
    buf = io.BytesIO()
    Image.new("RGB", size, (200 - i * 40, 120, 90)).save(buf, "JPEG", quality=95)
    imgs.append((f"x{i}.jpg", buf.getvalue()))

r = images_to_pdf(imgs, str(OUT / "_i.pdf"), page_size="A4", margin_mm=10)
check("ได้ครบ 3 หน้า", r.pages == 3, f"ได้ {r.pages}")
pdf = fitz.open(OUT / "_i.pdf")
check("หน้าแรกเป็นแนวนอนตามรูป", pdf[0].rect.width > pdf[0].rect.height)
check("หน้าที่สองเป็นแนวตั้งตามรูป", pdf[1].rect.height > pdf[1].rect.width)
pdf.close()

r2 = images_to_pdf(imgs, str(OUT / "_i2.pdf"), page_size="fit", quality_mode="original")
pdf2 = fitz.open(OUT / "_i2.pdf")
check("โหมดเท่าต้นฉบับ ขนาดหน้าตรงตามสัดส่วนรูป",
      abs(pdf2[0].rect.width / pdf2[0].rect.height - 1200 / 900) < 0.01)
pdf2.close()

print("\n=== 9. ไฟล์ที่เปิดไม่ได้ต้องแจ้งเตือน ไม่ใช่พัง ===")
bad = OUT / "_bad.pdf"
bad.write_bytes(b"not a pdf at all")
try:
    extract_document(str(bad))
    check("โยนข้อผิดพลาดเมื่อไฟล์เสีย", False, "ไม่โยนอะไรเลย")
except Exception:
    check("โยนข้อผิดพลาดเมื่อไฟล์เสีย", True)

try:
    images_to_pdf([("x.jpg", b"garbage")], str(OUT / "_none.pdf"))
    check("โยนข้อผิดพลาดเมื่อรูปเสียทั้งหมด", False)
except ValueError:
    check("โยนข้อผิดพลาดเมื่อรูปเสียทั้งหมด", True)

print("\n" + "=" * 60)
print(f"ผ่าน {PASS} ตก {FAIL}")
print("=" * 60)
sys.exit(1 if FAIL else 0)
