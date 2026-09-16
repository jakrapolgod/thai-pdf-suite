# -*- coding: utf-8 -*-
"""
build_docx.py - สร้างไฟล์ Word จากเอกสารที่ดึงมา

กับดักสำคัญของภาษาไทยใน Word:
Word จัดภาษาไทยอยู่ในกลุ่ม "complex script" ซึ่งใช้ค่าฟอนต์คนละชุดกับอักษรละติน
ถ้าตั้งแค่ w:ascii / w:hAnsi แต่ไม่ตั้ง w:cs  Word จะเลือกฟอนต์เองสำหรับอักษรไทย
ผลคือวรรณยุกต์ลอยผิดตำแหน่งหรือสระซ้อนกัน ทั้งที่ code point ถูกต้องทุกตัว

โมดูลนี้จึงตั้งครบทั้งชุด: w:cs (ฟอนต์), w:szCs (ขนาด), w:bCs/w:iCs (หนา/เอียง)
และประกาศภาษา th-TH ให้ตัวตรวจคำสะกดของ Word ด้วย
"""
from __future__ import annotations

from typing import List, Optional

from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .extract import Document, Paragraph, Run, Table

ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

# ฟอนต์ที่รองรับวรรณยุกต์ไทยได้ครบ เรียงตามลำดับความนิยมในเอกสารไทย
THAI_FONT_FALLBACK = ["TH Sarabun New", "Sarabun", "Angsana New", "Leelawadee UI", "Tahoma"]


def _set_run_font(run, name: str, size: float, bold: bool, italic: bool,
                  color: Optional[int]):
    """ตั้งค่าฟอนต์ให้ครบทั้งฝั่งละตินและฝั่ง complex script"""
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None and color != 0:
        run.font.color.rgb = RGBColor((color >> 16) & 0xFF, (color >> 8) & 0xFF, color & 0xFF)

    rpr = run._element.get_or_add_rPr()

    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), name)

    # ขนาดฝั่ง complex script ต้องตั้งแยก มิฉะนั้น Word ใช้ค่าเริ่มต้น 10pt กับอักษรไทย
    szcs = rpr.makeelement(qn("w:szCs"), {})
    szcs.set(qn("w:val"), str(int(round(size * 2))))
    rpr.append(szcs)

    if bold:
        rpr.append(rpr.makeelement(qn("w:bCs"), {}))
    if italic:
        rpr.append(rpr.makeelement(qn("w:iCs"), {}))

    lang = rpr.makeelement(qn("w:lang"), {})
    lang.set(qn("w:bidi"), "th-TH")
    rpr.append(lang)


def _apply_default_style(doc, font_name: str, size: float):
    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), font_name)
    szcs = rpr.makeelement(qn("w:szCs"), {})
    szcs.set(qn("w:val"), str(int(round(size * 2))))
    rpr.append(szcs)


def _add_paragraph(doc, para: Paragraph, font_name: str, scale: float):
    p = doc.add_paragraph()
    p.alignment = ALIGN_MAP.get(para.align, WD_ALIGN_PARAGRAPH.LEFT)
    pf = p.paragraph_format
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)

    for run in para.runs:
        if not run.text:
            continue
        r = p.add_run(run.text)
        _set_run_font(r, font_name, max(run.size * scale, 6.0),
                      run.bold, run.italic, run.color)
    return p


def _add_table(doc, table: Table, font_name: str, base_size: float):
    if not table.rows:
        return
    cols = max(len(r) for r in table.rows)
    t = doc.add_table(rows=len(table.rows), cols=cols)
    t.style = "Table Grid"

    for ri, row in enumerate(table.rows):
        for ci in range(cols):
            value = row[ci] if ci < len(row) else ""
            cell = t.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(value)
            _set_run_font(r, font_name, base_size, ri == 0, False, 0)
    doc.add_paragraph()


def build_docx(document: Document, out_path: str, *,
               font_name: str = "TH Sarabun New",
               font_scale: float = 1.0,
               keep_page_size: bool = True,
               page_break_between_pages: bool = True,
               base_size: float = 12.0) -> str:
    """
    สร้างไฟล์ .docx จากเอกสารที่ดึงมา
    font_scale ใช้ชดเชยกรณีเปลี่ยนไปใช้ฟอนต์ที่ตัวอักษรเล็กกว่าเดิม เช่น TH Sarabun
    """
    doc = DocxDocument()
    _apply_default_style(doc, font_name, base_size)

    if keep_page_size and document.pages:
        first = document.pages[0]
        section = doc.sections[0]
        if first.width > first.height:
            section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Pt(first.width)
        section.page_height = Pt(first.height)
        section.left_margin = Pt(56)
        section.right_margin = Pt(56)
        section.top_margin = Pt(56)
        section.bottom_margin = Pt(56)

    for idx, page in enumerate(document.pages):
        if idx > 0 and page_break_between_pages:
            doc.add_page_break()

        # เรียงย่อหน้าและตารางตามลำดับที่ปรากฏบนหน้ากระดาษ
        items = [("p", p.bbox[1], p) for p in page.paragraphs]
        items += [("t", t.bbox[1], t) for t in page.tables]
        items.sort(key=lambda it: it[1])

        for kind, _, obj in items:
            if kind == "p":
                _add_paragraph(doc, obj, font_name, font_scale)
            else:
                _add_table(doc, obj, font_name, base_size)

    doc.save(out_path)
    return out_path


def read_docx_text(path: str) -> str:
    """อ่านข้อความทั้งหมดกลับจากไฟล์ Word สำหรับใช้ตรวจสอบความตรงกัน"""
    doc = DocxDocument(path)
    parts: List[str] = []

    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            text = "".join(node.text or "" for node in child.iter(qn("w:t")))
            if text.strip():
                parts.append(text)
        elif tag == "tbl":
            for tr in child.iter(qn("w:tr")):
                cells = []
                for tc in tr.iter(qn("w:tc")):
                    cells.append("".join(node.text or "" for node in tc.iter(qn("w:t"))))
                if any(c.strip() for c in cells):
                    parts.append("\t".join(cells))
    return "\n".join(parts)
